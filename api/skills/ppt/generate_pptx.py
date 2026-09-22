"""PPT 可编辑产物生成（指南 v18 P1-2，ppt-master/dashi-ppt 对标）。

输入：ppt 技能大纲 JSON（title/pages[headline, points, notes]/assumptions）
输出：可编辑 OOXML PPTX bytes（python-pptx 渲染，标题 + 要点 + 演讲备注）。
零浏览器依赖；模板体系（P2）后置。纯函数，可单测锁定（文件头 PK + slide 数）。
"""
from __future__ import annotations

import io
from typing import Any


def generate_pptx(outline: dict[str, Any]) -> bytes:
    """大纲 → PPTX bytes。outline 缺省字段容错（不崩，最小可用闭环）。"""
    from pptx import Presentation

    prs = Presentation()
    title = str(outline.get("title") or "演示文稿").strip() or "演示文稿"
    pages = outline.get("pages") or []

    # 封面页（大标题）
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    if slide.shapes.title:
        slide.shapes.title.text = title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = str(outline.get("subtitle") or "")

    # 正文页：一页一观点（headline 即结论 + 3-5 要点 + 备注）
    for i, page in enumerate(pages or [], start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        if slide.shapes.title:
            headline = str(page.get("headline") or f"第 {i} 页").strip() or f"第 {i} 页"
            slide.shapes.title.text = headline
        body = slide.placeholders[1].text_frame if len(slide.placeholders) > 1 else None
        if body is not None:
            first = True
            for point in (page.get("points") or [])[:5]:
                p = body.paragraphs[0] if first else body.add_paragraph()
                first = False
                p.text = str(point)
        notes = str(page.get("notes") or "")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
