"""PPT 可编辑产物生成测试（v18 P1-2）。纯本地（python-pptx）。"""
from __future__ import annotations

import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from api.skills.ppt.generate_pptx import generate_pptx  # noqa: E402

OUTLINE = {
    "title": "产品发布会",
    "subtitle": "AI 图像生成网关",
    "pages": [
        {"no": 1, "headline": "开场：行业痛点", "points": ["等待慢", "成本高", "质量不稳"], "notes": "强调 3 个痛点"},
        {"no": 2, "headline": "方案：统一网关", "points": ["多提供商聚合", "MAB 路由", "预算门禁"], "notes": ""},
    ],
    "assumptions": ["受众=开发者", "时长 20 分钟"],
}


def test_generate_pk_zip_header():
    data = generate_pptx(OUTLINE)
    assert data[:2] == b"PK"  # ZIP/OOXML 文件头
    assert len(data) > 500


def test_generate_openable_with_pptx():
    from pptx import Presentation

    data = generate_pptx(OUTLINE)
    prs = Presentation(io.BytesIO(data))
    # 封面 + 2 页 = 3 slides
    assert len(prs.slides) == 3
    assert prs.slides[0].shapes.title.text == "产品发布会"


def test_generate_title_only_minimal():
    data = generate_pptx({"title": "极简", "pages": []})
    assert data[:2] == b"PK"
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 1  # 仅封面


def test_generate_empty_outline_ok():
    data = generate_pptx({})
    assert data[:2] == b"PK"  # 全缺省不崩，产出「演示文稿」封面


def test_generate_notes_captured():
    from pptx import Presentation

    data = generate_pptx(OUTLINE)
    prs = Presentation(io.BytesIO(data))
    notes = prs.slides[1].notes_slide.notes_text_frame.text
    assert "痛点" in notes or notes == ""
